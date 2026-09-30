import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def create_element(name):
    return OxmlElement(name)

def set_cell_background(cell, fill_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="CCCCCC", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:left w:val="none"/>
            <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:right w:val="none"/>
            <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:insideV w:val="none"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)

def build_document(output_path):
    doc = Document()

    # Page Margins (1 inch = 1440 dxa)
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Styles & Colors
    NAVY = RGBColor(0, 51, 102)      # #003366
    DARK_GRAY = RGBColor(51, 51, 51) # #333333
    BLUE_BG = "003366"
    GRAY_BG = "F2F4F7"
    BORDER_COLOR = "CCCCCC"

    # --- PORTADA ---
    p_institution = doc.add_paragraph()
    p_institution.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run_inst = p_institution.add_run("DUOC UC  |  MALETA DIDÁCTICA\nEscuela de Informática y Telecomunicaciones")
    run_inst.font.name = "Arial"
    run_inst.font.size = Pt(10)
    run_inst.font.bold = True
    run_inst.font.color.rgb = NAVY

    doc.add_paragraph().paragraph_format.space_after = Pt(40)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r_title = p_title.add_run("ACTIVIDAD FORMATIVA 2\n")
    r_title.font.name = "Arial"
    r_title.font.size = Pt(24)
    r_title.font.bold = True
    r_title.font.color.rgb = NAVY

    r_subtitle = p_title.add_run("Creando Pruebas Unitarias y de Inspección de Código")
    r_subtitle.font.name = "Arial"
    r_subtitle.font.size = Pt(16)
    r_subtitle.font.color.rgb = DARK_GRAY

    doc.add_paragraph().paragraph_format.space_after = Pt(80)

    # Meta Table in Portada
    meta_table = doc.add_table(rows=5, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False

    meta_data = [
        ("Asignatura:", "Técnicas de Calidad de Software (TCY0101)"),
        ("Proyecto:", "Tienda MoAr (React + TypeScript + Vitest + ESLint)"),
        ("Integrantes / Roles:", "Nicolas Iturrieta (Analista QA)\nMiguel Arredondo (Jefe de Proyecto QA)"),
        ("Profesor:", "Christian Álvarez Lobos"),
        ("Fecha de Entrega:", "Septiembre 2026")
    ]

    for i, (label, val) in enumerate(meta_data):
        row = meta_table.rows[i]
        c1, c2 = row.cells[0], row.cells[1]
        c1.width = Inches(2.2)
        c2.width = Inches(4.3)
        
        p1 = c1.paragraphs[0]
        r1 = p1.add_run(label)
        r1.font.bold = True
        r1.font.name = "Arial"
        r1.font.size = Pt(10)
        r1.font.color.rgb = NAVY
        
        p2 = c2.paragraphs[0]
        r2 = p2.add_run(val)
        r2.font.name = "Arial"
        r2.font.size = Pt(10)
        r2.font.color.rgb = DARK_GRAY

        set_cell_background(c1, "F8F9FA")
        set_cell_margins(c1, top=120, bottom=120, left=150, right=150)
        set_cell_margins(c2, top=120, bottom=120, left=150, right=150)

    set_table_borders(meta_table, color="D0D5DD", sz="4")

    doc.add_page_break()

    # --- SECCIÓN 1: INTRODUCCIÓN Y CONTEXTO ---
    h1 = doc.add_heading(level=1)
    r_h1 = h1.add_run("1. Instrucciones Generales y Contexto")
    r_h1.font.name = "Arial"
    r_h1.font.color.rgb = NAVY

    p_intro = doc.add_paragraph()
    p_intro.paragraph_format.line_spacing = 1.15
    p_intro.paragraph_format.space_after = Pt(10)
    p_intro.add_run(
        "Esta actividad tiene como finalidad configurar pruebas Unitarias y de Inspección de Código dentro de la aplicación web "
        "Tienda MoAr, apoyado en el material de la sesión y la guía del docente. El trabajo fue realizado en dupla, asumiendo los roles de "
        "Analista QA (Nicolas Iturrieta) y Jefe de Proyecto QA (Miguel Arredondo).\n\n"
        "El sistema Tienda MoAr es una plataforma de e-commerce tecnológica desarrollada en React, TypeScript y Vite. "
        "En el marco del Plan de Pruebas de la Parcial No. 1, se definió como punto crítico el cálculo automático de los subtotales "
        "y el total general del carrito de compras (Requerimiento Funcional RF14 / Caso de Prueba CP19)."
    )

    # --- SECCIÓN 2: PARTE I - PRUEBAS UNITARIAS ---
    h2 = doc.add_heading(level=1)
    r_h2 = h2.add_run("2. Parte I: Pruebas Unitarias (Caja Blanca)")
    r_h2.font.name = "Arial"
    r_h2.font.color.rgb = NAVY

    p_p1 = doc.add_paragraph()
    p_p1.paragraph_format.line_spacing = 1.15
    p_p1.paragraph_format.space_after = Pt(10)
    p_p1.add_run(
        "Se configuró e instaló el framework de pruebas unitarias Vitest (v4.0.14), el cual se integra de forma nativa con Vite y TypeScript. "
        "Se seleccionaron 4 guiones de prueba derivados del caso de prueba CP19 (RF14-1) documentado en el Plan de Pruebas Parcial 1, "
        "evaluando la función de negocio calcularSubtotal del módulo src/utils/cart.ts."
    )

    # Tabla I
    p_tbl1_title = doc.add_paragraph()
    r_t1 = p_tbl1_title.add_run("Tabla I: Registro y Resultados de Pruebas Unitarias (Caja Blanca)")
    r_t1.font.bold = True
    r_t1.font.color.rgb = NAVY

    t1 = doc.add_table(rows=1, cols=7)
    t1.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t1, color="B0C4DE", sz="4")

    headers1 = ["Id", "Caso de Prueba", "Descripción", "Datos / Acciones de Entrada", "Resultado Esperado", "Resultado Obtenido", "Observaciones y Evidencia"]
    col_widths1 = [Inches(0.6), Inches(1.1), Inches(1.3), Inches(1.3), Inches(0.8), Inches(0.8), Inches(1.1)]

    hdr_cells1 = t1.rows[0].cells
    for i, title in enumerate(headers1):
        hdr_cells1[i].width = col_widths1[i]
        set_cell_background(hdr_cells1[i], BLUE_BG)
        set_cell_margins(hdr_cells1[i], top=120, bottom=120, left=100, right=100)
        p = hdr_cells1[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(title)
        r.font.bold = True
        r.font.name = "Arial"
        r.font.size = Pt(8.5)
        r.font.color.rgb = RGBColor(255, 255, 255)

    data1 = [
        ("CP19-1", "Carrito Vacío", "Verificar cálculo con arreglo vacío (RF14-1).", "lines = []", "0", "0", "Aprobado (0ms) - ✓ Vitest"),
        ("CP19-2", "Línea Única", "Validar multiplicación precio * cantidad (RF14-1).", "lines = [{precioUnitario: 1000, cantidad: 3}]", "3000", "3000", "Aprobado (1ms) - ✓ Vitest"),
        ("CP19-3", "Múltiples Líneas", "Validar suma acumulativa de varios ítems (RF14-1).", "lines = [{precio: 1000, cant: 2}, {precio: 500, cant: 4}]", "4000", "4000", "Aprobado (1ms) - ✓ Vitest"),
        ("CP19-4", "Precios Elevados", "Verificar cálculo con valores reales de smartphones.", "lines = [{precio: 1299990, cant: 1}, {precio: 299990, cant: 2}]", "1899970", "1899970", "Aprobado (1ms) - ✓ Vitest")
    ]

    for row_idx, row_data in enumerate(data1):
        row_cells = t1.add_row().cells
        bg = GRAY_BG if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, cell_value in enumerate(row_data):
            row_cells[col_idx].width = col_widths1[col_idx]
            set_cell_background(row_cells[col_idx], bg)
            set_cell_margins(row_cells[col_idx], top=100, bottom=100, left=80, right=80)
            p = row_cells[col_idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if col_idx in [0, 4, 5] else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(cell_value)
            r.font.name = "Arial"
            r.font.size = Pt(8)
            r.font.color.rgb = DARK_GRAY

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # Console output block
    p_code_title = doc.add_paragraph()
    r_ct = p_code_title.add_run("Evidencia de Ejecución en Consola (Vitest Terminal Output):")
    r_ct.font.bold = True
    r_ct.font.size = Pt(9.5)
    r_ct.font.color.rgb = NAVY

    p_code = doc.add_paragraph()
    p_code.paragraph_format.space_after = Pt(15)
    run_code = p_code.add_run(
        " RUN  v4.0.14 C:/Users/xkait/Documents/MoAr/tienda-moar\n\n"
        " ✓ src/tests/cart.test.ts (4 tests) 5ms\n\n"
        " Test Files  1 passed (1)\n"
        "      Tests  4 passed (4)\n"
        "   Duration  2.58s"
    )
    run_code.font.name = "Consolas"
    run_code.font.size = Pt(8.5)
    run_code.font.color.rgb = RGBColor(0, 102, 0)

    # --- SECCIÓN 3: PARTE II - INSPECCIÓN DE CÓDIGO ---
    h3 = doc.add_heading(level=1)
    r_h3 = h3.add_run("3. Parte II: Pruebas de Inspección de Código (Análisis Estático)")
    r_h3.font.name = "Arial"
    r_h3.font.color.rgb = NAVY

    p_p2 = doc.add_paragraph()
    p_p2.paragraph_format.line_spacing = 1.15
    p_p2.paragraph_format.space_after = Pt(10)
    p_p2.add_run(
        "Se configuró la herramienta de inspección estática ESLint v9 con los plugins de TypeScript y React. "
        "Se ejecutó un escaneo completo sobre los componentes, controladores y páginas del sistema WEB Tienda MoAr (src/). "
        "A continuación se presenta la síntesis del análisis estático realizado:"
    )

    # Tabla II
    p_tbl2_title = doc.add_paragraph()
    r_t2 = p_tbl2_title.add_run("Tabla II: Análisis de Resultados del Escaneo Estático")
    r_t2.font.bold = True
    r_t2.font.color.rgb = NAVY

    t2 = doc.add_table(rows=1, cols=3)
    t2.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t2, color="B0C4DE", sz="4")

    headers2 = ["Categoría", "Cantidad", "Breve Análisis y Diagnóstico"]
    col_widths2 = [Inches(1.8), Inches(1.0), Inches(4.2)]

    hdr_cells2 = t2.rows[0].cells
    for i, title in enumerate(headers2):
        hdr_cells2[i].width = col_widths2[i]
        set_cell_background(hdr_cells2[i], BLUE_BG)
        set_cell_margins(hdr_cells2[i], top=120, bottom=120, left=100, right=100)
        p = hdr_cells2[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(title)
        r.font.bold = True
        r.font.name = "Arial"
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(255, 255, 255)

    data2 = [
        ("Bugs", "0", "No se detectaron fallos sintácticos o bloqueos críticos de runtime en el código escaneado."),
        ("Vulnerabilidades", "0", "No se detectaron patrones de inyección ni riesgos de seguridad directos en el frontend."),
        ("Code Smells", "10", "Inconsistencias de mantenibilidad: uso de tipo 'any' implícito/explícito (Carrito.tsx, RegistroCliente.tsx), variables declaradas sin uso ('subtotal' y 'Product'), y hooks exportados junto con contextos React."),
        ("Cobertura", "100%", "Cobertura total del suite de pruebas unitarias sobre el módulo de cálculo de subtotales de compra (cart.ts).")
    ]

    for row_idx, (cat, cant, anal) in enumerate(data2):
        row_cells = t2.add_row().cells
        bg = GRAY_BG if row_idx % 2 == 1 else "FFFFFF"
        
        row_cells[0].width = col_widths2[0]
        set_cell_background(row_cells[0], bg)
        set_cell_margins(row_cells[0], top=100, bottom=100, left=100, right=100)
        p0 = row_cells[0].paragraphs[0]
        r0 = p0.add_run(cat)
        r0.font.bold = True
        r0.font.name = "Arial"
        r0.font.size = Pt(8.5)

        row_cells[1].width = col_widths2[1]
        set_cell_background(row_cells[1], bg)
        set_cell_margins(row_cells[1], top=100, bottom=100, left=100, right=100)
        p1 = row_cells[1].paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r1 = p1.add_run(cant)
        r1.font.bold = True
        r1.font.name = "Arial"
        r1.font.size = Pt(8.5)
        if cant in ["0"]:
            r1.font.color.rgb = RGBColor(0, 128, 0)
        elif cant == "10":
            r1.font.color.rgb = RGBColor(180, 100, 0)

        row_cells[2].width = col_widths2[2]
        set_cell_background(row_cells[2], bg)
        set_cell_margins(row_cells[2], top=100, bottom=100, left=100, right=100)
        p2 = row_cells[2].paragraphs[0]
        r2 = p2.add_run(anal)
        r2.font.name = "Arial"
        r2.font.size = Pt(8.5)
        r2.font.color.rgb = DARK_GRAY

    doc.add_paragraph().paragraph_format.space_after = Pt(15)

    # --- SECCIÓN 4: PARTE III - INFORME Y CONCLUSIONES ---
    h4 = doc.add_heading(level=1)
    r_h4 = h4.add_run("4. Parte III: Informe de Procedimiento y Conclusiones")
    r_h4.font.name = "Arial"
    r_h4.font.color.rgb = NAVY

    conclusions = [
        ("Procedimiento Realizado: ", "Se inició con la auditoría de dependencias y la puesta en marcha de Vitest como runner de pruebas unitarias de caja blanca. Posteriormente, se alinearon los casos de prueba de la función de cálculo del carrito con los requerimientos del Plan de Pruebas (RF14/CP19) redactado para la Parcial 1."),
        ("Resultados de Pruebas Unitarias: ", "La función calcularSubtotal respondió satisfactoriamente bajo los 4 escenarios planteados (arreglo vacío, ítem único, múltiples ítems y precios elevados de smartphones), alcanzando un 100% de tasa de éxito."),
        ("Resultados de Inspección de Código: ", "El escaneo con ESLint identificó 10 observaciones leves de mantenibilidad (Code Smells). Se definió un plan de acción para reemplazar el tipado 'any' por interfaces estrictas de TypeScript y remover variables inactivas en la siguiente iteración."),
        ("Criterios de Aceptación: ", "El software cumple con los criterios de aceptación en términos de integridad de cálculo de compras y ausencia de bugs críticos, estando apto para su validación en fase de QA.")
    ]

    for title, body in conclusions:
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_after = Pt(6)
        r_t = p.add_run(title)
        r_t.font.bold = True
        r_t.font.name = "Arial"
        r_t.font.size = Pt(10)
        r_t.font.color.rgb = NAVY

        r_b = p.add_run(body)
        r_b.font.name = "Arial"
        r_b.font.size = Pt(10)
        r_b.font.color.rgb = DARK_GRAY

    doc.add_paragraph().paragraph_format.space_after = Pt(30)

    # Signatures Table
    sig_table = doc.add_table(rows=1, cols=2)
    sig_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_s1, c_s2 = sig_table.rows[0].cells
    c_s1.width = Inches(3.2)
    c_s2.width = Inches(3.2)

    p_s1 = c_s1.paragraphs[0]
    p_s1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_s1.add_run("_____________________________________\nNicolas Iturrieta\nAnalista QA").font.name = "Arial"

    p_s2 = c_s2.paragraphs[0]
    p_s2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_s2.add_run("_____________________________________\nMiguel Arredondo\nJefe de Proyecto QA").font.name = "Arial"

    doc.save(output_path)
    print("Document created successfully at:", output_path)

if __name__ == "__main__":
    build_document("Unidad 2/Actividad_Formativa_2.2.3_TiendaMoAr.docx")
