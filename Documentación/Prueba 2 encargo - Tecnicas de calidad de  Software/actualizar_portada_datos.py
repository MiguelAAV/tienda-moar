# -*- coding: utf-8 -*-
"""
Script para actualizar los datos de la portada en Evidencias_Presentacion_Final.pptx
Agrega los datos de asignatura, sigla, profesor e integrantes manteniendo el diseño y el resto de diapositivas intacto.
"""

import os
import glob
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def main():
    files = glob.glob('**/Evidencias_Presentacion_Final.pptx', recursive=True)
    if not files:
        print("No se encontró el archivo de presentación.")
        return
    
    target_file = files[0]
    prs = pptx.Presentation(target_file)
    s1 = prs.slides[0]

    # Eliminar formas previas que no sean las imágenes o conectores originales del usuario (primeras 5 formas)
    # Conservamos exactamente las formas 0 a 4 agregadas por el usuario
    while len(s1.shapes) > 5:
        shape_to_remove = s1.shapes[len(s1.shapes) - 1]
        sp = shape_to_remove._element
        sp.getparent().remove(sp)

    # 1. Línea divisoria vertical simétrica en x=9.10"
    line = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(9.10), Inches(1.72), Inches(0.015), Inches(4.06))
    line.fill.solid()
    line.fill.fore_color.rgb = RGBColor(127, 127, 127) # #7F7F7F
    line.line.color.rgb = RGBColor(127, 127, 127)
    line.line.width = Pt(1)

    # 2. Tarjeta con datos informativos en la columna derecha
    tb = s1.shapes.add_textbox(Inches(9.35), Inches(1.25), Inches(3.60), Inches(5.10))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

    # Etiqueta de la evaluación
    p0 = tf.paragraphs[0]
    p0.text = "EVALUACIÓN PARCIAL 2.2 (25%)"
    p0.font.name = "Arial"
    p0.font.size = Pt(10)
    p0.font.bold = True
    p0.font.color.rgb = RGBColor(37, 99, 235) # Azul Royal
    p0.space_after = Pt(4)

    # Título principal del sistema
    p1 = tf.add_paragraph()
    p1.text = "Tienda MoAr"
    p1.font.name = "Arial"
    p1.font.size = Pt(28)
    p1.font.bold = True
    p1.font.color.rgb = RGBColor(15, 23, 42) # Azul Marino Profundo
    p1.space_after = Pt(2)

    # Subtítulo del proyecto
    p2 = tf.add_paragraph()
    p2.text = "Aseguramiento de Calidad y Pruebas"
    p2.font.name = "Arial"
    p2.font.size = Pt(12)
    p2.font.color.rgb = RGBColor(71, 85, 105) # Slate
    p2.space_after = Pt(18)

    # Lista de metadatos con espaciado y tipografía nítida
    metadata = [
        ("Asignatura:", "Técnicas de Calidad de Software"),
        ("Sigla:", "TCY0101"),
        ("Docente:", "Christian Álvarez Lobos"),
        ("Integrantes:", "• Miguel Arredondo\n• Nicolás Iturrieta"),
        ("Fecha:", "Octubre 2026")
    ]

    for label, val in metadata:
        p = tf.add_paragraph()
        
        r_lbl = p.add_run()
        r_lbl.text = label + "\n" if "\n" in val else label + " "
        r_lbl.font.bold = True
        r_lbl.font.size = Pt(11)
        r_lbl.font.name = "Arial"
        r_lbl.font.color.rgb = RGBColor(30, 41, 59)
        
        r_val = p.add_run()
        r_val.text = val
        r_val.font.bold = False
        r_val.font.size = Pt(11)
        r_val.font.name = "Arial"
        r_val.font.color.rgb = RGBColor(71, 85, 105)
        
        p.space_after = Pt(8)

    # Notas de orador para la diapositiva 1
    notes_slide = s1.notes_slide
    tf_notes = notes_slide.notes_text_frame
    tf_notes.text = (
        "[Tiempo: ~20 seg | Hablan: Miguel Arredondo y Nicolás Iturrieta]\n"
        "Buenas tardes profesor Christian Álvarez. Somos Miguel Arredondo y Nicolás Iturrieta, "
        "y a continuación presentamos la Evaluación Parcial 2.2 de Técnicas de Calidad de Software "
        "para nuestra plataforma e-commerce Tienda MoAr."
    )

    prs.save(target_file)
    print("Portada actualizada correctamente con acentos y formato impecable.")

if __name__ == "__main__":
    main()
